from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any
import re

import pandas as pd

from src.utils import format_currency, format_percent


# -----------------------------------------------------------------------------
# Excel helpers
# -----------------------------------------------------------------------------

def _flatten_tables(obj: Any, prefix: str = "") -> dict[str, pd.DataFrame]:
    tables: dict[str, pd.DataFrame] = {}
    if isinstance(obj, pd.DataFrame):
        if not obj.empty:
            tables[prefix[:31] or "Table"] = obj
        return tables
    if isinstance(obj, dict):
        for key, value in obj.items():
            name = f"{prefix}_{key}" if prefix else str(key)
            tables.update(_flatten_tables(value, name))
    elif isinstance(obj, list) and obj and isinstance(obj[0], dict):
        tables[prefix[:31] or "Items"] = pd.DataFrame(obj)
    return tables


def _safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = "".join(ch for ch in name.replace(".", "_").replace("/", "_") if ch not in "[]:*?\\")[:31] or "Sheet"
    base = cleaned
    idx = 1
    while cleaned in used:
        suffix = f"_{idx}"
        cleaned = (base[: 31 - len(suffix)] + suffix)
        idx += 1
    used.add(cleaned)
    return cleaned


def _style_workbook(writer) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="0B1F3A")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00'
        for column_cells in ws.columns:
            letter = get_column_letter(column_cells[0].column)
            max_len = max(len(str(c.value)) if c.value is not None else 0 for c in column_cells)
            ws.column_dimensions[letter].width = min(55, max(12, max_len + 2))


def _summary_rows(context_results: dict[str, Any]) -> pd.DataFrame:
    summary = context_results.get("executive_summary", {}) or {}
    kpis = context_results.get("kpis", {}) or {}
    missing_meta = context_results.get("missing_value_treatment", {}) or {}
    rows = []
    if missing_meta.get("cleaned_dataset_used"):
        rows.append({"Section": "Data Reliability Note", "Finding": missing_meta.get("manager_warning")})
        rows.append({"Section": "Missing Value Treatment Summary", "Finding": missing_meta.get("summary_text")})
    for key, value in summary.items():
        rows.append({"Section": key.replace("_", " ").title(), "Finding": value})
    if kpis:
        rows.extend([
            {"Section": "Total Sales", "Finding": format_currency(kpis.get("total_sales"))},
            {"Section": "Total Profit", "Finding": format_currency(kpis.get("total_profit"))},
            {"Section": "Profit Margin", "Finding": format_percent(kpis.get("profit_margin_percent"))},
            {"Section": "Loss-Making Orders", "Finding": kpis.get("loss_making_orders_count")},
        ])
    return pd.DataFrame(rows)


def _missing_treatment_df(context_results: dict[str, Any]) -> pd.DataFrame:
    meta = context_results.get("missing_value_treatment", {}) if isinstance(context_results, dict) else {}
    rows = meta.get("treatment_log", []) if isinstance(meta, dict) else []
    return pd.DataFrame(rows) if rows else pd.DataFrame()


def generate_excel_report(df, column_roles, context_results):
    """Generate a formatted multi-sheet Excel report as bytes."""
    output = BytesIO()
    used: set[str] = set()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _summary_rows(context_results).to_excel(writer, index=False, sheet_name=_safe_sheet_name("Executive Summary", used))
        pd.DataFrame([context_results.get("kpis", {})]).to_excel(writer, index=False, sheet_name=_safe_sheet_name("KPIs", used))
        treatment = _missing_treatment_df(context_results)
        if treatment.empty:
            treatment = pd.DataFrame([{"Status": "No missing-value treatment was required."}])
        treatment.to_excel(writer, index=False, sheet_name=_safe_sheet_name("Missing Value Treatment", used))
        recs = context_results.get("recommendations", [])
        if recs:
            pd.DataFrame(recs).to_excel(writer, index=False, sheet_name=_safe_sheet_name("Recommendations", used))

        exact_sheets = {
            "Profit Intelligence": context_results.get("profit_intelligence", {}).get("profit_leakage_points"),
            "Customer Intelligence": context_results.get("customer_intelligence", {}).get("top_loyal_customers"),
            "Product Intelligence": context_results.get("product_intelligence", {}).get("problem_products"),
            "Region Intelligence": context_results.get("region_intelligence", {}).get("region_summary"),
            "Forecasting": context_results.get("forecasting", {}).get("forecast"),
            "Anomalies": context_results.get("anomaly_detection", {}).get("unusual_rows"),
        }
        for sheet_name, table in exact_sheets.items():
            if isinstance(table, pd.DataFrame) and not table.empty:
                table.head(5000).to_excel(writer, index=False, sheet_name=_safe_sheet_name(sheet_name, used))
            else:
                pd.DataFrame([{"Status": "No data available for this section."}]).to_excel(writer, index=False, sheet_name=_safe_sheet_name(sheet_name, used))

        tables = _flatten_tables(context_results)
        preferred = [
            "profit_intelligence_profit_leakage_points",
            "metric_relationships_profit_by_discount_band",
            "metric_relationships_high_sales_low_profit",
            "customer_intelligence_top_loyal_customers",
            "customer_intelligence_unprofitable_customers",
            "product_intelligence_problem_products",
            "region_intelligence_weak_regions",
            "forecasting_forecast",
            "anomaly_detection_unusual_rows",
        ]
        for key in preferred + [k for k in tables if k not in preferred]:
            table = tables.get(key)
            if isinstance(table, pd.DataFrame) and not table.empty:
                safe = _safe_sheet_name(key, used)
                table.head(5000).to_excel(writer, index=False, sheet_name=safe)
        df.head(5000).to_excel(writer, index=False, sheet_name=_safe_sheet_name("Raw Data Sample", used))
        _style_workbook(writer)
    output.seek(0)
    return output.getvalue()


# -----------------------------------------------------------------------------
# PDF helpers
# -----------------------------------------------------------------------------

PDF_LABELS = {
    "business_performance": "Business Performance",
    "key_problem": "Key Problem",
    "root_cause": "Root Cause",
    "financial_impact": "Financial Impact",
    "recommended_action": "Recommended Action",
    "risk": "Risk",
    "next_step": "Next Step",
}


def _plain_money(value: Any) -> str:
    """Format money for PDF output using USD currency."""
    return format_currency(value)


def _clean_text(value: Any) -> str:
    """Convert backend text into PDF-safe manager text."""
    if value is None:
        return "N/A"
    text = str(value)
    text = text.replace("₹", "$")
    text = text.replace("\u2022", "-")
    text = text.replace("•", "-")
    text = re.sub(r"\s+", " ", text).strip()
    text = text.replace(";.", ".")
    text = text.replace(";;", ";")
    return text


def _clean_items(items: Any, limit: int | None = None) -> list[str]:
    if items is None:
        return []
    if isinstance(items, str):
        raw = re.split(r";\s*|\n+", items)
    elif isinstance(items, list):
        raw = items
    else:
        raw = [items]
    out: list[str] = []
    seen: set[str] = set()
    for item in raw:
        text = _clean_text(item).strip(" .;-")
        if not text:
            continue
        key = text.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(text + ".")
        if limit and len(out) >= limit:
            break
    return out


def _df(obj: Any) -> pd.DataFrame:
    return obj if isinstance(obj, pd.DataFrame) else pd.DataFrame()


def _norm_col_name(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(name).lower())


def _pick_columns(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame()
    lookup = {_norm_col_name(c): c for c in df.columns}
    picked = []
    for requested in columns:
        actual = lookup.get(_norm_col_name(requested))
        if actual and actual not in picked:
            picked.append(actual)
    if not picked:
        return pd.DataFrame()
    return df[picked].copy()


def _format_cell(value: Any, col: str = "") -> str:
    if pd.isna(value) if not isinstance(value, (list, dict, pd.DataFrame)) else False:
        return ""
    col_l = str(col).lower()
    try:
        if any(term in col_l for term in ["margin", "percent", "%", "discount", "rate"]):
            v = float(value)
            # Discount values may be stored as 0.45. Percent values may already be 45.
            if "discount" in col_l and abs(v) <= 1:
                v *= 100
            return f"{v:.2f}%"
        if any(term in col_l for term in ["sales", "profit", "loss", "cost", "amount", "aov"]):
            return _plain_money(value)
        if isinstance(value, float):
            return f"{value:,.2f}"
        if isinstance(value, int):
            return f"{value:,}"
    except Exception:
        pass
    return _clean_text(value)


def _title_case(label: str) -> str:
    return str(label).replace("_", " ").replace("dimension value", "Name").title()


def _table_data_from_df(df: pd.DataFrame, columns: list[str] | None = None, max_rows: int = 8) -> list[list[str]]:
    if df is None or df.empty:
        return []
    work = _pick_columns(df, columns) if columns else df.copy()
    if work.empty:
        return []
    work = work.head(max_rows).copy()
    header = [_title_case(c) for c in work.columns]
    rows = []
    for _, row in work.iterrows():
        rows.append([_format_cell(row[c], c) for c in work.columns])
    return [header] + rows


def _para(text: Any, style):
    from xml.sax.saxutils import escape
    from reportlab.platypus import Paragraph
    return Paragraph(escape(_clean_text(text)), style)


def _bold_para(label: str, text: Any, style):
    from xml.sax.saxutils import escape
    from reportlab.platypus import Paragraph
    return Paragraph(f"<b>{escape(_clean_text(label))}:</b> {escape(_clean_text(text))}", style)


def _bullet_list(items: list[str], style, limit: int = 6) -> list[Any]:
    from reportlab.platypus import Paragraph, Spacer
    from xml.sax.saxutils import escape
    story: list[Any] = []
    for idx, item in enumerate(_clean_items(items, limit=limit), start=1):
        story.append(Paragraph(f"<b>{idx}.</b> {escape(item)}", style))
        story.append(Spacer(1, 4))
    return story


def _styled_table(data: list[list[Any]], widths: list[float] | None = None, small: bool = False, header_color: str = "#0B1F3A"):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet

    if not data:
        return None
    styles = getSampleStyleSheet()
    body_style = styles["BodyText"]
    body_style.fontSize = 7 if small else 8
    body_style.leading = 9 if small else 10
    cell_data = []
    for r_idx, row in enumerate(data):
        cell_row = []
        for cell in row:
            cell_row.append(Paragraph(_clean_text(cell), body_style))
        cell_data.append(cell_row)
    table = Table(cell_data, hAlign="LEFT", colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(header_color)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("FONTSIZE", (0, 1), (-1, -1), 7 if small else 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return table


def _section(story: list[Any], title: str, styles, subtitle: str | None = None) -> None:
    from reportlab.platypus import Paragraph, Spacer, HRFlowable
    from reportlab.lib import colors
    story.append(Spacer(1, 10))
    story.append(Paragraph(title, styles["SectionTitle"]))
    if subtitle:
        story.append(Paragraph(subtitle, styles["Muted"]))
    story.append(HRFlowable(width="100%", thickness=0.7, color=colors.HexColor("#D7E3F2"), spaceBefore=3, spaceAfter=8))


def _page_footer(canvas, doc) -> None:
    from reportlab.lib import colors
    width, height = doc.pagesize
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D7E3F2"))
    canvas.line(doc.leftMargin, 28, width - doc.rightMargin, 28)
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawString(doc.leftMargin, 18, "AI Sales Analyst Copilot - Manager Intelligence Report")
    canvas.drawRightString(width - doc.rightMargin, 18, f"Page {doc.page}")
    canvas.restoreState()


def _kpi_card_table(kpis: dict[str, Any]):
    from reportlab.lib.units import inch
    cards = [
        ["Total Sales", _plain_money(kpis.get("total_sales"))],
        ["Total Profit", _plain_money(kpis.get("total_profit"))],
        ["Profit Margin", format_percent(kpis.get("profit_margin_percent"))],
        ["Orders", _format_cell(kpis.get("total_orders"), "orders")],
        ["Loss Orders", _format_cell(kpis.get("loss_making_orders_count"), "orders")],
        ["Avg Discount", _format_cell(kpis.get("average_discount"), "discount")],
    ]
    data = []
    for i in range(0, len(cards), 3):
        labels = [cards[i + j][0] if i + j < len(cards) else "" for j in range(3)]
        values = [cards[i + j][1] if i + j < len(cards) else "" for j in range(3)]
        data.append(labels)
        data.append(values)
    table = _styled_table(data, widths=[2.2 * inch, 2.2 * inch, 2.2 * inch], header_color="#123A5A")
    if table:
        from reportlab.lib import colors
        from reportlab.platypus import TableStyle
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EAFDF8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0B1F3A")),
            ("TEXTCOLOR", (0, 2), (-1, 2), colors.HexColor("#0B1F3A")),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 1), (-1, 1), 13),
            ("FONTSIZE", (0, 3), (-1, 3), 13),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#99F6E4")),
        ]))
    return table


def _build_recommendation_table(recs: list[dict[str, Any]], max_rows: int = 6) -> list[list[str]]:
    data = [["Priority", "Action", "Evidence", "Expected Impact"]]
    for rec in recs[:max_rows]:
        data.append([
            rec.get("priority", ""),
            rec.get("recommended_action", rec.get("title", "")),
            rec.get("evidence_from_data", rec.get("issue_detected", "")),
            rec.get("expected_impact", ""),
        ])
    return data


def generate_pdf_report(context_results):
    """Generate an information-rich manager-ready PDF report as bytes.

    The report is designed for leadership review: summary first, then evidence,
    action priorities, root-cause details, and appendix-style diagnostic tables.
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, KeepTogether

    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=34,
        leftMargin=34,
        topMargin=34,
        bottomMargin=42,
        title="AI Sales Analyst Copilot Manager Report",
        author="AI Sales Analyst Copilot",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], fontSize=23, leading=28, textColor=colors.HexColor("#0B1F3A"), spaceAfter=6))
    styles.add(ParagraphStyle(name="ReportSubtitle", parent=styles["Heading2"], fontSize=12, leading=16, textColor=colors.HexColor("#0F766E"), spaceAfter=12))
    styles.add(ParagraphStyle(name="SectionTitle", parent=styles["Heading2"], fontSize=14, leading=18, textColor=colors.HexColor("#123A5A"), spaceBefore=8, spaceAfter=4))
    styles.add(ParagraphStyle(name="SubTitle", parent=styles["Heading3"], fontSize=10, leading=13, textColor=colors.HexColor("#0F766E"), spaceBefore=7, spaceAfter=4))
    styles.add(ParagraphStyle(name="Body", parent=styles["BodyText"], fontSize=8.5, leading=11.5, textColor=colors.HexColor("#0F172A"), spaceAfter=4))
    styles.add(ParagraphStyle(name="Muted", parent=styles["BodyText"], fontSize=8, leading=10.5, textColor=colors.HexColor("#64748B"), spaceAfter=5))
    styles.add(ParagraphStyle(name="Callout", parent=styles["BodyText"], fontSize=9, leading=12, textColor=colors.HexColor("#0B1F3A"), backColor=colors.HexColor("#EAFDF8"), borderColor=colors.HexColor("#99F6E4"), borderWidth=0.6, borderPadding=7, spaceAfter=8))

    story: list[Any] = []
    summary = context_results.get("executive_summary", {}) or {}
    kpis = context_results.get("kpis", {}) or {}
    pi = context_results.get("profit_intelligence", {}) or {}
    rel = context_results.get("metric_relationships", {}) or {}
    dq = context_results.get("data_quality", {}) or {}
    cust = context_results.get("customer_intelligence", {}) or {}
    prod = context_results.get("product_intelligence", {}) or {}
    reg = context_results.get("region_intelligence", {}) or {}
    forecast = context_results.get("forecasting", {}) or {}
    anomalies = context_results.get("anomaly_detection", {}) or {}
    hidden = context_results.get("hidden_insights", []) or []
    recs = context_results.get("recommendations", []) or []
    missing_meta = context_results.get("missing_value_treatment", {}) or {}

    # Cover / executive header
    story.append(Paragraph("AI Sales Analyst Copilot", styles["ReportTitle"]))
    story.append(Paragraph("Manager-Ready Business Intelligence Report", styles["ReportSubtitle"]))
    story.append(Paragraph(f"Generated on {datetime.now().strftime('%d %b %Y, %I:%M %p')}. All numeric findings are calculated from the uploaded dataset.", styles["Muted"]))
    if isinstance(missing_meta, dict) and missing_meta.get("cleaned_dataset_used"):
        story.append(Paragraph(_clean_text(missing_meta.get("manager_warning")), styles["Callout"]))
        if missing_meta.get("summary_text"):
            story.append(Paragraph(_clean_text(missing_meta.get("summary_text")), styles["Muted"]))
    story.append(Spacer(1, 8))

    if kpis:
        cards = _kpi_card_table(kpis)
        if cards:
            story.append(cards)
            story.append(Spacer(1, 10))

    if summary:
        story.append(Paragraph("Executive Decision Brief", styles["SectionTitle"]))
        brief_data = [["Area", "Manager Brief"]]
        for key in ["business_performance", "key_problem", "root_cause", "financial_impact", "recommended_action", "risk", "next_step"]:
            if summary.get(key):
                brief_data.append([PDF_LABELS.get(key, _title_case(key)), summary[key]])
        story.append(_styled_table(brief_data, widths=[1.55 * inch, 5.1 * inch], small=False, header_color="#0B1F3A"))

    # What management should do first
    if recs:
        _section(story, "Priority Actions", styles, "Ranked actions based on financial impact, confidence, and visible root causes.")
        story.append(_styled_table(_build_recommendation_table(recs, 7), widths=[0.65 * inch, 2.05 * inch, 2.25 * inch, 1.75 * inch], small=True, header_color="#0F766E"))

    # Profit intelligence
    _section(story, "Profit Intelligence", styles, "Explains margin leakage rather than only reporting top-line sales.")
    root_causes = _clean_items(pi.get("root_cause_summary"), limit=8)
    if root_causes:
        story.append(Paragraph("Key Drivers", styles["SubTitle"]))
        story.extend(_bullet_list(root_causes, styles["Body"], limit=8))
    relationship_items = _clean_items(rel.get("relationship_summary"), limit=8)
    if relationship_items:
        story.append(Paragraph("Metric Relationships", styles["SubTitle"]))
        story.extend(_bullet_list(relationship_items, styles["Body"], limit=8))

    discount_df = _df(rel.get("profit_by_discount_band"))
    if not discount_df.empty:
        story.append(Paragraph("Profit by Discount Band", styles["SubTitle"]))
        data = _table_data_from_df(discount_df, ["Discount Band", "sales", "profit", "loss_amount", "orders", "profit_margin_percent"], 8)
        if data:
            story.append(_styled_table(data, widths=[1.1 * inch, 1.1 * inch, 1.1 * inch, 1.1 * inch, 0.75 * inch, 1.05 * inch], small=True, header_color="#123A5A"))

    leakage_df = _df(pi.get("profit_leakage_points"))
    if not leakage_df.empty:
        story.append(Paragraph("Top Profit Leakage Points", styles["SubTitle"]))
        data = _table_data_from_df(leakage_df, ["dimension", "dimension_value", "sales", "profit", "profit_margin_percent", "loss_amount", "order_count"], 8)
        if data:
            story.append(_styled_table(data, widths=[0.9 * inch, 1.35 * inch, 1.0 * inch, 1.0 * inch, 0.9 * inch, 0.95 * inch, 0.65 * inch], small=True, header_color="#123A5A"))

    # Product, customer, region evidence
    story.append(PageBreak())
    _section(story, "Product, Customer, and Regional Evidence", styles, "Operational drilldowns behind the executive summary.")

    problem_products = _df(prod.get("problem_products"))
    if not problem_products.empty:
        story.append(Paragraph("Problem Products", styles["SubTitle"]))
        data = _table_data_from_df(problem_products, ["dimension_value", "sales", "profit", "profit_margin_percent", "average_discount", "order_count"], 8)
        if data:
            story.append(_styled_table(data, widths=[1.75 * inch, 1.1 * inch, 1.1 * inch, 0.95 * inch, 0.95 * inch, 0.7 * inch], small=True, header_color="#B45309"))

    portfolio = _df(prod.get("product_portfolio"))
    if not portfolio.empty:
        story.append(Paragraph("Portfolio Snapshot", styles["SubTitle"]))
        data = _table_data_from_df(portfolio, ["dimension_value", "portfolio_segment", "sales", "profit", "profit_margin_percent"], 8)
        if data:
            story.append(_styled_table(data, widths=[1.8 * inch, 1.45 * inch, 1.1 * inch, 1.1 * inch, 1.0 * inch], small=True, header_color="#0F766E"))

    loyal_customers = _df(cust.get("top_loyal_customers"))
    if not loyal_customers.empty:
        story.append(Paragraph("Top Loyal Customers", styles["SubTitle"]))
        data = _table_data_from_df(loyal_customers, ["customer", "order_count", "total_sales", "total_profit", "profit_margin_percent", "loyalty_score"], 8)
        if data:
            story.append(_styled_table(data, widths=[1.6 * inch, 0.75 * inch, 1.05 * inch, 1.05 * inch, 0.95 * inch, 0.9 * inch], small=True, header_color="#0F766E"))

    weak_regions = _df(reg.get("weak_regions"))
    if not weak_regions.empty:
        story.append(Paragraph("Weak Regions", styles["SubTitle"]))
        data = _table_data_from_df(weak_regions, ["dimension_value", "sales", "profit", "profit_margin_percent", "average_discount", "order_count"], 8)
        if data:
            story.append(_styled_table(data, widths=[1.35 * inch, 1.1 * inch, 1.1 * inch, 1.0 * inch, 1.0 * inch, 0.75 * inch], small=True, header_color="#B45309"))

    # Hidden insights, forecast, data quality
    story.append(PageBreak())
    _section(story, "Hidden Insights, Forecasts, and Trust Checks", styles, "Signals that management should validate before changing operations.")

    if hidden:
        story.append(Paragraph("Hidden Insights", styles["SubTitle"]))
        insight_data = [["Severity", "Insight", "Why It Matters", "Recommended Action"]]
        for item in hidden[:7]:
            if isinstance(item, dict):
                insight_data.append([
                    item.get("severity", ""),
                    item.get("title", ""),
                    item.get("why_it_matters", item.get("evidence", "")),
                    item.get("recommended_action", item.get("action", "")),
                ])
            else:
                insight_data.append(["", str(item), "", ""])
        story.append(_styled_table(insight_data, widths=[0.7 * inch, 1.7 * inch, 2.1 * inch, 2.1 * inch], small=True, header_color="#123A5A"))

    fdf = _df(forecast.get("forecast"))
    if not fdf.empty:
        story.append(Paragraph("Forecast Preview", styles["SubTitle"]))
        data = _table_data_from_df(fdf, None, 6)
        if data:
            story.append(_styled_table(data, widths=None, small=True, header_color="#0F766E"))
        if forecast.get("limitations"):
            story.append(Paragraph("Forecast Limitations", styles["SubTitle"]))
            story.extend(_bullet_list(forecast.get("limitations", []), styles["Body"], limit=4))

    anomaly_items = anomalies.get("anomalies", []) if isinstance(anomalies, dict) else []
    if anomaly_items:
        story.append(Paragraph("Anomalies to Review", styles["SubTitle"]))
        anomaly_data = [["Severity", "Metric", "Explanation", "Recommended Check"]]
        for item in anomaly_items[:7]:
            anomaly_data.append([item.get("severity", ""), item.get("metric", ""), item.get("explanation", item.get("title", "")), item.get("recommended_check", "")])
        story.append(_styled_table(anomaly_data, widths=[0.75 * inch, 1.0 * inch, 2.75 * inch, 2.1 * inch], small=True, header_color="#B45309"))

    if dq:
        story.append(Paragraph("Data Quality and Business Signal Separation", styles["SubTitle"]))
        dq_rows = [["Check", "Result"]]
        for label, key in [("Quality Score", "quality_score"), ("Duplicate Rows", "duplicate_rows_count"), ("Missing Cells %", "missing_cells_percent")]:
            if key in dq:
                value = format_percent(dq[key]) if "percent" in key else _format_cell(dq[key], key)
                dq_rows.append([label, value])
        for signal in dq.get("business_signals", [])[:4] if isinstance(dq.get("business_signals"), list) else []:
            dq_rows.append(["Business Signal", signal])
        story.append(_styled_table(dq_rows, widths=[1.8 * inch, 4.8 * inch], small=True, header_color="#0B1F3A"))
        treatment_df = _missing_treatment_df(context_results)
        if not treatment_df.empty:
            story.append(Paragraph("Missing Value Treatment Log", styles["SubTitle"]))
            data = _table_data_from_df(
                treatment_df,
                ["column", "role", "missing_before", "treatment", "values_fixed", "values_still_missing", "business_impact"],
                10,
            )
            if data:
                story.append(_styled_table(data, widths=[0.85 * inch, 0.75 * inch, 0.7 * inch, 1.7 * inch, 0.65 * inch, 0.7 * inch, 1.35 * inch], small=True, header_color="#0F766E"))
    else:
        story.append(Paragraph("Data quality details were not included in this report context. Use the app's Data Quality tab for trust checks.", styles["Body"]))

    # Management action plan
    _section(story, "Suggested 30-Day Management Plan", styles)
    plan_items = [
        "Week 1: Validate the top loss driver with finance, sales, and operations owners.",
        "Week 2: Audit discount policy and cost assumptions for the top loss-making product/category/region.",
        "Week 3: Run a controlled pricing or discount-cap test on the affected segment.",
        "Week 4: Track profit margin, loss order percentage, and customer retention impact before scaling the change.",
    ]
    story.extend(_bullet_list(plan_items, styles["Body"], limit=4))

    story.append(Paragraph("Limitations", styles["SectionTitle"]))
    story.append(Paragraph(
        "This report is generated from detected columns in the uploaded dataset. The system separates technical data quality from business-performance signals such as negative profit. Forecasts and scenarios are estimates, not guarantees. Validate recommendations with business owners before operational changes.",
        styles["Body"],
    ))

    doc.build(story, onFirstPage=_page_footer, onLaterPages=_page_footer)
    output.seek(0)
    return output.getvalue()
